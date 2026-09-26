"""Build soft-magnet-paper-v3.docx from manuscript_v3.txt.

The 30 Zotero citation fields and the cached bibliography of the original
"Summer Paper.docx" are copied as intact XML subtrees; their instruction streams
are never rewritten.  New references that were not in the original library are
added as plain red [xN] markers and a red reference list, so the Zotero field
codes remain the single source of truth for the existing citations.
"""
from pathlib import Path
from copy import deepcopy
from collections import Counter
from zipfile import ZipFile
import json
import re
import hashlib

import numpy as np
import pandas as pd
from lxml import etree
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

W = Path(__file__).resolve().parent
F = W / 'figures_v3'
SRC_DOC = Path('/mnt/user-data/uploads/soft-magnets/paper/Summer Paper.docx')
DST = W / 'soft-magnet-paper-v3.docx'
NS = {'w': 'http://schemas.openxmlformats.org/wordprocessingml/2006/main'}

source_hash = hashlib.sha256(SRC_DOC.read_bytes()).hexdigest()
doc = Document(SRC_DOC)
orig = list(doc._element.body)

fields = []
for block in orig[:59]:
    for p in ([block] if block.tag == qn('w:p') else block.findall('.//w:p', NS)):
        active, depth = [], 0
        for node in p:
            types = [s.get(qn('w:fldCharType')) for s in node.findall('.//w:fldChar', NS)]
            if 'begin' in types:
                depth += 1
            if depth:
                active.append(deepcopy(node))
            if 'end' in types:
                depth -= 1
                if depth == 0:
                    fields.append(active)
                    active = []
assert len(fields) == 30, len(fields)
bib = [deepcopy(b) for b in orig[60:90]]

for child in list(doc._element.body):
    if child.tag != qn('w:sectPr'):
        doc._element.body.remove(child)

sec = doc.sections[0]
sec.page_width, sec.page_height = Inches(8.5), Inches(11)
sec.top_margin = sec.bottom_margin = Inches(.78)
sec.left_margin = sec.right_margin = Inches(.9)
sec.header_distance, sec.footer_distance = Inches(.3), Inches(.32)

for name in ['Normal', 'Body Text', 'Title', 'Subtitle', 'Heading 1', 'Heading 2',
             'Heading 3', 'Caption']:
    if name not in doc.styles:
        doc.styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH)
    st = doc.styles[name]
    st.font.name = 'Times New Roman'
    st.font.color.rgb = RGBColor(0, 0, 0)
    st.paragraph_format.space_before = Pt(0)
    st.paragraph_format.space_after = Pt(6)
    st.paragraph_format.line_spacing = 1.08
doc.styles['Normal'].font.size = Pt(11)
doc.styles['Normal'].paragraph_format.first_line_indent = Inches(0)
for name, size in [('Title', 16.5), ('Heading 1', 13), ('Heading 2', 11.5),
                   ('Heading 3', 11), ('Caption', 9.5)]:
    doc.styles[name].font.size = Pt(size)
for name in ['Heading 1', 'Heading 2', 'Heading 3']:
    doc.styles[name].font.bold = True
    doc.styles[name].paragraph_format.space_before = Pt(12)
    doc.styles[name].paragraph_format.keep_with_next = True
doc.styles['Title'].font.bold = True
doc.styles['Title'].paragraph_format.space_after = Pt(10)
doc.styles['Caption'].font.italic = False
doc.styles['Caption'].paragraph_format.space_after = Pt(9)

# Prevent automatic field refresh; do not alter Zotero instructions or cached values.
up = doc.settings.element.find(qn('w:updateFields'))
if up is not None:
    doc.settings.element.remove(up)

footer = sec.footer.paragraphs[0]
footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
for c in list(footer._p):
    footer._p.remove(c)
r = footer.add_run()
fld = OxmlElement('w:fldSimple')
fld.set(qn('w:instr'), 'PAGE')
r._r.addnext(fld)

used = []
SUBSCRIPT = (r'(T_C|T_c|H_c|M_s|λ_s|B_0|K_1|K_2|J_ij|k_i|x_i|v_i|p_i|ρ_i|α_i|y_i|ŷ_i'
             r'|r_j|q_1−α|μ_B|E_σ→0|E_FCC|N_FCC|E_BCC|N_BCC|T_0\.1|E_\[111\]|E_\[001\])')


def add_plain(p, text, red=False):
    """Human-readable subscripts in prose; displayed equations use native OMML."""
    for s in re.split(SUBSCRIPT, text):
        if not s:
            continue
        if re.fullmatch(SUBSCRIPT, s):
            a, b = s.split('_', 1)
            run = p.add_run(a)
            run.font.color.rgb = RGBColor(255, 0, 0) if red else RGBColor(0, 0, 0)
            run = p.add_run(b)
            run.font.subscript = True
            if red:
                run.font.color.rgb = RGBColor(255, 0, 0)
        else:
            run = p.add_run(s)
            if red:
                run.font.color.rgb = RGBColor(255, 0, 0)


def rich(p, text):
    for s in re.split(r'(\[c\d+\]|\[x\d+\])', text):
        if re.fullmatch(r'\[c\d+\]', s):
            i = int(s[2:-1])
            used.append(i)
            for n in fields[i]:
                p._p.append(deepcopy(n))
        elif re.fullmatch(r'\[x\d+\]', s):
            add_plain(p, s, True)
        else:
            add_plain(p, s)
    return p


def para(text, style=None):
    p = doc.add_paragraph(style=style)
    rich(p, text)
    p.paragraph_format.widow_control = True
    if style is None:
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    return p


def append_block(el):
    doc._element.body.insert(len(doc._element.body) - 1, deepcopy(el))


# ---- native Word math ------------------------------------------------------
def el(name, *children):
    e = OxmlElement('m:' + name)
    for c in children:
        e.append(c)
    return e


def mr(text):
    t = el('t')
    t.text = text
    return el('r', t)


def seq(*parts):
    return [mr(p) if isinstance(p, str) else p for p in parts]


def box(name, parts):
    return el(name, *seq(*parts))


def sub(base, index):
    return el('sSub', box('e', [base]), box('sub', [index]))


def sup(base, power):
    return el('sSup', box('e', [base]), box('sup', [power]))


def frac(a, b):
    return el('f', box('num', a), box('den', b))


def sm(expr, low='i=1', up='n'):
    pr = el('naryPr')
    ch = el('chr'); ch.set(qn('m:val'), '∑'); pr.append(ch)
    loc = el('limLoc'); loc.set(qn('m:val'), 'undOvr'); pr.append(loc)
    return el('nary', pr, box('sub', [low]), box('sup', [up]), box('e', expr))


def root(a):
    pr = el('radPr')
    h = el('degHide'); h.set(qn('m:val'), '1'); pr.append(h)
    return el('rad', pr, el('deg'), box('e', a))


def dbox(parts):
    return el('d', box('e', parts))


def eq(parts, num):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(5)
    p.paragraph_format.space_after = Pt(9)
    math = el('oMath', *seq(*parts))
    p._p.append(math)
    p.add_run('    (' + str(num) + ')').font.size = Pt(10.5)


def equations(n):
    yi, yh = sub('y', 'i'), sub('ŷ', 'i')
    if n == 1:
        eq(['VEC = ', sm([sub('x', 'i'), sub('v', 'i')], up='p'), ',     ',
            sub('ρ', 'desc'), ' = ', sm([sub('x', 'i'), sub('ρ', 'i')], up='p')], 1)
    if n == 2:
        eq([sub('p', 'mean'), ' = ', sm([sub('x', 'i'), sub('p', 'i')], up='p'),
            ',     ', sub('p', 'absdev'), ' = ',
            sm([sub('x', 'i'), '|', sub('p', 'i'), ' − ', sub('p', 'mean'), '|'], up='p')], '2a')
        eq([sub('p', 'max'), ' = ', 'max ', sub('p', 'i'), ',     ',
            sub('p', 'min'), ' = ', 'min ', sub('p', 'i'), ',     ',
            sub('p', 'range'), ' = ', sub('p', 'max'), ' − ', sub('p', 'min')], '2b')
    if n == 3:
        eq([sub('y', 'H'), ' = ', sub('log', '10'), '(',
            frac([sub('H', 'c')], ['1 Oe']), ')'], 3)
    if n == 4:
        eq([sup('R', '2'), ' = 1 − ',
            frac([sm([sup(dbox([yi, ' − ', yh]), '2')])],
                 [sm([sup(dbox([yi, ' − ȳ']), '2')])])], '4a')
        eq(['MAE = ', frac(['1'], ['n']), sm(['|', yi, ' − ', yh, '|']), ',     RMSE = ',
            root([frac(['1'], ['n']), sm([sup(dbox([yi, ' − ', yh]), '2')])])], '4b')
    if n == 5:
        eq([sub('q', '1−α'), ' = ', sub('Q', 'β'), dbox(['{', sub('r', 'j'), '}']),
            ',     β = ', frac(['⌈(m + 1)(1 − α)⌉'], ['m']),
            ',     C(x) = ŷ(x) ± ', sub('q', '1−α')], 5)
    if n == 6:
        eq(['Δe(x) = ', frac([sub('E', 'FCC'), '(x)'], [sub('N', 'FCC')]), ' − ',
            frac([sub('E', 'BCC'), '(x)'], [sub('N', 'BCC')])], 6)
    if n == 7:
        a = lambda i: sup(sub('α', str(i)), '2')
        eq([sub('e', 'ani'), ' = ', sub('K', '1'), '(', a(1), a(2), ' + ', a(2), a(3),
            ' + ', a(3), a(1), ') + ', sub('K', '2'), a(1), a(2), a(3)], 7)
    if n == 8:
        eq(['ℋ = −', sm([sub('J', 'ij'), sub('s', 'i'), ' · ', sub('s', 'j')],
                        low='i<j', up=''), ' − ',
            sm([sub('k', 'i'), sup(dbox([sub('s', 'i'), ' · ', sub('n', 'i')]), '2')],
               low='i', up='')], 8)
    if n == 9:
        eq([sub('J', 'ij'), ' = ',
            frac([sub('E', '+−'), ' + ', sub('E', '−+'), ' − ', sub('E', '++'),
                  ' − ', sub('E', '−−')], ['4'])], 9)


# ---- tables ----------------------------------------------------------------
def table(headers, rows, widths, fontsize=9):
    t = doc.add_table(rows=1, cols=len(headers))
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    for c, w in zip(t.columns, widths):
        c.width = Inches(w)
    for i, h in enumerate(headers):
        rich(t.rows[0].cells[i].paragraphs[0], h)
    for row in rows:
        cells = t.add_row().cells
        for c, txt in zip(cells, row):
            rich(c.paragraphs[0], str(txt))
    pr = t._tbl.tblPr
    borders = OxmlElement('w:tblBorders')
    for name in ['top', 'left', 'bottom', 'right', 'insideH', 'insideV']:
        b = OxmlElement('w:' + name)
        b.set(qn('w:val'), 'single'); b.set(qn('w:sz'), '4'); b.set(qn('w:color'), 'D9D9D9')
        borders.append(b)
    pr.append(borders)
    for ri, row in enumerate(t.rows):
        trpr = row._tr.get_or_add_trPr()
        trpr.append(OxmlElement('w:cantSplit'))
        if ri == 0:
            trpr.append(OxmlElement('w:tblHeader'))
        for ci, c in enumerate(row.cells):
            c.width = Inches(widths[ci])
            c.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            tcpr = c._tc.get_or_add_tcPr()
            mar = OxmlElement('w:tcMar')
            for name, val in [('top', 60), ('bottom', 60), ('left', 70), ('right', 70)]:
                z = OxmlElement('w:' + name)
                z.set(qn('w:w'), str(val)); z.set(qn('w:type'), 'dxa')
                mar.append(z)
            tcpr.append(mar)
            if ri == 0:
                sh = OxmlElement('w:shd'); sh.set(qn('w:fill'), 'E8EDF2'); tcpr.append(sh)
            for p in c.paragraphs:
                p.paragraph_format.space_after = Pt(0)
                p.paragraph_format.line_spacing = 1.0
                p.paragraph_format.keep_with_next = False
                for run in p.runs:
                    # Existing Zotero field nodes remain entirely untouched.
                    if (run._r.find(qn('w:instrText')) is None
                            and not run._r.findall('.//w:fldChar', NS)):
                        run.font.size = Pt(fontsize)
                        run.bold = (ri == 0)
    return t


def caption(s):
    return para(s, 'Caption')


def tables(name):
    if name == 'metrics':
        caption('Table 1. Held-out performance by partitioning protocol. The first two '
                'protocols share encoded compositions between fitting and testing. The '
                'grouped repeats report the mean and standard deviation over twenty '
                'independent partitions with nested selection of descriptor set and '
                'learner; the pooled value is computed over all held-out predictions '
                'together. Coercivity errors refer to log₁₀[H_c/(1 Oe)].')
        old = pd.read_csv(F / 'ml_metrics.csv')
        ns = pd.read_csv(W / 'v3_nested_summary.csv')
        pooled = pd.read_csv(W / 'v3_pooled_predictions.csv')
        from sklearn.metrics import r2_score, mean_absolute_error
        rows = []
        names = {'notebook_fixed': 'Notebook random', 'composition_random': 'Composition random'}
        for _, r in old[old.protocol != 'composition_grouped'].iterrows():
            tl = 'T_C (K)' if r.target.startswith('Curie') else 'log₁₀ H_c'
            rows.append([names[r.protocol], tl, int(r.n_test), int(r.test_composition_overlap),
                         f'{r.r2:.3f}', f'{r.mae:.3f}', f'{r.rmse:.3f}'])
        for _, r in old[old.protocol == 'composition_grouped'].iterrows():
            tl = 'T_C (K)' if r.target.startswith('Curie') else 'log₁₀ H_c'
            rows.append(['Grouped, single partition', tl, int(r.n_test), 0,
                         f'{r.r2:.3f}', f'{r.mae:.3f}', f'{r.rmse:.3f}'])
        nm = pd.read_csv(W / 'v3_nested_metrics.csv')
        nm = nm[nm.split != 'v2_reference_split']
        ns = ns.set_index('target')
        for t in ['Curie(TC) (K)', 'Coercivity (Oe)']:
            r = ns.loc[t]
            tl = 'T_C (K)' if t.startswith('Curie') else 'log₁₀ H_c'
            g = nm[nm.target == t]
            rows.append(['Grouped, 20 repeats', tl,
                         f'{g.n_test.mean():.0f} (mean)', 0,
                         f'{r.r2_mean:.3f} ± {r.r2_sd:.3f}',
                         f'{r.mae_mean:.3f} ± {r.mae_sd:.3f}', f'{r.rmse_mean:.3f}'])
        for t in ['Curie(TC) (K)', 'Coercivity (Oe)']:
            g = pooled[pooled.target == t]
            tl = 'T_C (K)' if t.startswith('Curie') else 'log₁₀ H_c'
            rows.append(['Grouped, pooled predictions', tl, len(g), 0,
                         f'{r2_score(g.observed, g.predicted):.3f}',
                         f'{mean_absolute_error(g.observed, g.predicted):.3f}', '—'])
        table(['Protocol', 'Target', 'n test', 'Shared compositions', 'R²', 'MAE', 'RMSE'],
              rows, [1.70, .78, .60, .90, 1.18, 1.05, .72], fontsize=8.5)

    if name == 'descriptors':
        caption('Table 2. Held-out R² by descriptor set on identical composition-grouped '
                'partitions, as mean ± standard deviation over the repeats, together with '
                'how often the nested protocol selected each set. Weighted elemental '
                'statistics help Curie temperature and do not help coercivity.')
        d = pd.read_csv(F / 'figure4_descriptors.csv')
        sel = pd.read_csv(W / 'v3_nested_metrics.csv')
        sel = sel[sel.split != 'v2_reference_split']
        labels = {'fractions': 'Elemental fractions (79)',
                  'physics': 'Weighted elemental statistics (146)',
                  'combined': 'Both (225)'}
        rows = []
        for fs in ['fractions', 'physics', 'combined']:
            for t, tl in [('Curie(TC) (K)', 'T_C (K)'), ('Coercivity (Oe)', 'log₁₀ H_c')]:
                g = d[(d.feature_set == fs) & (d.target == t)].iloc[0]
                n = int((sel[(sel.target == t)].selected_feature_set == fs).sum())
                rows.append([labels[fs], tl, f'{g.r2_mean:.3f} ± {g.r2_sd:.3f}',
                             f'{g.mae_mean:.3f}', f'{n} / 20'])
        table(['Descriptor set (features)', 'Target', 'Held-out R²', 'MAE', 'Times selected'],
              rows, [2.30, .82, 1.32, .88, 1.05], fontsize=8.8)

    if name == 'candidates':
        caption('Table 3. The eight highest-ranked supported recipes at the resolvable '
                'target T_C > 600 K and H_c < 100 Oe, from 147 that satisfy it. Bounds are '
                'the calibrated 90% conformal limits. Fractions are atomic percent. The '
                'ranking does not resolve differences smaller than the interval width, and '
                'no entry is a validated material.')
        d = pd.read_csv(F / 'supported_recipes_600K_100Oe.csv').head(8)
        rows = []
        for i, (_, r) in enumerate(d.iterrows(), 1):
            comp = '  '.join(f'{e} {100*r[e]:.1f}' for e in
                             ['Fe', 'Co', 'Ni', 'Mn', 'Al', 'Si'])
            rows.append([f'R{i}', comp, f'{r.predicted_Tc_K:.0f} ({r.Tc_lower_K:.0f})',
                         f'{r.predicted_Hc_Oe:.1f} ({r.Hc_upper_Oe:.0f})',
                         f'{r.L1_to_nearest_training:.3f}',
                         'interp.' if r.novelty == 'interpolation' else 'near-extrap.'])
        table(['ID', 'Composition (at.%)', 'T_C (K), lower bound',
               'H_c (Oe), upper bound', 'L1 distance', 'Regime'],
              rows, [.38, 2.62, 1.18, 1.14, .78, .85], fontsize=8.2)

    if name == 'benchmark':
        caption('Table 4. Evaluation protocols used in recent machine-learning studies of '
                'soft magnetic and magnetic alloys. Reported scores are not commensurate: '
                'they differ in target property, data provenance and partitioning. The '
                'column that can be compared is whether chemically identical material was '
                'prevented from appearing on both sides of the split.')
        rows = [
            ['Wang et al. 2020 [x1]', 'Fe-based nanocrystalline', 'H_c, M_s, λ_s',
             'experiment', '20-fold CV, random', 'No', 'reported graphically'],
            ['Milyutin et al. 2024 [x11]', 'Fe–Si–Al ternary', 'H_c and others',
             'experiment', '85:15 random', 'No', '0.935 (H_c)'],
            ['Nachnani et al. 2025 [x2]', 'Fe-based soft magnets', 'M_s, H_c',
             'experiment', 'random with CV', 'No', '0.86 / 0.76'],
            ['Dai et al. 2025 [x3]', 'soft magnetic HEAs', 'phase, B₀, M, T_C',
             'DFT', '80:20 random', 'No', '0.921–0.987'],
            ['Abedi Orang et al. 2025 [x12]', 'general ferromagnets', 'T_C',
             'mixed compilation', '80:20 random, 5-fold CV', 'No', '0.66 / 0.87–0.91'],
            ['Meredig et al. 2018 [x10]', 'superconductors', 'T_c',
             'experiment', 'random vs leave-one-cluster-out', 'Yes (k-means clusters)',
             '0.87 vs 0.30'],
            ['This work', 'compiled magnetic materials', 'T_C, log₁₀ H_c',
             'experiment', 'composition-grouped, 20 repeats', 'Yes (exact compositions)',
             '0.601 ± 0.080 / 0.589 ± 0.069'],
        ]
        table(['Study', 'System', 'Target', 'Target source', 'Partition',
               'Composition leakage controlled', 'Reported R²'],
              rows, [1.42, 1.15, .80, .78, 1.28, 1.02, 1.20], fontsize=7.8)

    if name == 'exchange':
        caption('Table 5. Audited four-state combinations. Values are raw '
                '(E₊₋ + E₋₊ − E₊₊ − E₋₋)/4 and are not validated bond exchange constants. '
                'Context references concern different methods or structures and are not '
                'numerical validation targets.')
        d = pd.read_csv(F / 'exchange_audit.csv')
        notes = ['Nominal +− collapses to ++', 'One state reaches NELM',
                 'Nominal +− collapses to ++', 'Two states reach NELM',
                 'Incomplete first state; missing outputs',
                 'Missing OUTCARs; near-equal energies', 'Three missing OUTCARs',
                 'Missing OUTCARs; equal energies']
        lab = {'bccFe': 'bcc Fe', 'B2FeCo': 'B2 FeCo', 'L12Fe3Co': 'L1₂ Fe₃Co',
               'L12FeCo3': 'L1₂ FeCo₃'}
        pr = {'FeFe': 'Fe–Fe', 'FeCo': 'Fe–Co', 'CoCo': 'Co–Co'}
        rows = []
        for i, (_, r) in enumerate(d.iterrows()):
            mat, pair = r.run.split('_')[2], r.run.split('_')[-1]
            rows.append([lab[mat] + ' ' + pr[pair], f'{r.raw_combination_meV:.5f}',
                         notes[i], 'See text' if i == 0 else f'[c{22+i}]'])
        table(['System and pair', 'Raw value (meV)', 'Consistency assessment',
               'Context refs'], rows, [1.7, .88, 2.68, 1.22])


refs = [
 '[x1] Y. Wang, Y. Tian, T. Kirk, O. Laris, J. H. Ross Jr., R. D. Noebe, V. Keylin and R. Arróyave, Accelerated design of Fe-based soft magnetic materials using machine learning and stochastic optimization, Acta Materialia 194 (2020) 144–155. https://doi.org/10.1016/j.actamat.2020.05.006',
 '[x2] A. Nachnani, K. K. Li-Caldwell, S. Biswas, P. Sharma, G. Ouyang and P. Singh, Interpretable machine learning-guided design of Fe-based soft magnetic alloys, Physical Review Materials 9 (2025) 084411, and Supplemental Material. https://doi.org/10.1103/w6m3-ymsf',
 '[x3] M. Dai, Y. Zhang, X. Li, S. Schönecker, L. Han, R. Xie, C. Shen and H. Zhang, Data-Driven Design of Mechanically Hard Soft Magnetic High-Entropy Alloys, Advanced Science 12 (2025) 2500867. https://doi.org/10.1002/advs.202500867',
 '[x4] M. Dai, Y. Zhang, W. He, C. Shen, X. Li, S. Schönecker, L. Han, R. Xie, T. Zhou and H. Zhang, Accelerated Design of Mechanically Hard Magnetically Soft High-entropy Alloys via Multi-objective Bayesian Optimization, arXiv:2509.05702v1 (2025), preprint. https://doi.org/10.48550/arXiv.2509.05702',
 '[x5] Y. Yang, P. Kühn, M. Fathidoost, E. Adabifiroozjaei, R. Xie, E. Foya, D. Ohmer, K. Skokov, L. Molina-Luna, O. Gutfleisch, H. Zhang and B.-X. Xu, Coercivity influence of nanostructure in SmCo-1:7 magnets: machine learning of high-throughput micromagnetic data, npj Computational Materials 12 (2026) 204. https://doi.org/10.1038/s41524-026-02174-y',
 '[x6] Y. Tatetsu, K. Matsumoto, R. Sato and T. Teranishi, First-principles study of structural stability and magnetic properties in Fe–Pd–In alloys, Journal of Magnetism and Magnetic Materials 636 (2025) 173649. https://doi.org/10.1016/j.jmmm.2025.173649',
 '[x7] B. Balasubramanian, P. Manchanda, R. Skomski, P. Mukherjee, S. R. Valloppilly, B. Das, G. C. Hadjipanayis and D. J. Sellmyer, High-coercivity magnetism in nanostructures with strong easy-plane anisotropy, Applied Physics Letters 108 (2016) 152406. https://doi.org/10.1063/1.4945987',
 '[x8] VASP Software GmbH, SAXIS and ICHARG, VASP Wiki, accessed 25 September 2026. https://vasp.at/wiki/SAXIS ; https://vasp.at/wiki/ICHARG',
 '[x9] L. Ward, A. Agrawal, A. Choudhary and C. Wolverton, A general-purpose machine learning framework for predicting properties of inorganic materials, npj Computational Materials 2 (2016) 16028. https://doi.org/10.1038/npjcompumats.2016.28',
 '[x10] B. Meredig, E. Antono, C. Church, M. Hutchinson, J. Ling, S. Paradiso, B. Blaiszik, I. Foster, B. Gibbons, J. Hattrick-Simpers, A. Mehta and L. Ward, Can machine learning identify the next high-temperature superconductor? Examining extrapolation performance for materials discovery, Molecular Systems Design & Engineering 3 (2018) 819–825. https://doi.org/10.1039/C8ME00012C',
 '[x11] V. A. Milyutin, R. Bureš, M. Fáberová, Z. Birčáková, Z. Molčanová, B. Kunca, L. A. Stashkova, P. Kollár and J. Füzer, Machine learning assisted optimization of soft magnetic properties in ternary Fe–Si–Al alloys, Journal of Materials Research and Technology 29 (2024) 5060–5073. https://doi.org/10.1016/j.jmrt.2024.02.215',
 '[x12] A. Abedi Orang, M. Alaei and A. R. Oganov, Predicting the Curie Temperature of Magnetic Materials with Machine Learning: Descriptor Engineering, Graph Neural Networks, and the Role of Curated Data, arXiv:2509.17464 (2025), preprint. https://doi.org/10.48550/arXiv.2509.17464',
 '[x13] J. Lei, M. G’Sell, A. Rinaldo, R. J. Tibshirani and L. Wasserman, Distribution-Free Predictive Inference for Regression, Journal of the American Statistical Association 113 (2018) 1094–1111. https://doi.org/10.1080/01621459.2017.1307116',
]

for line in (W / 'manuscript_v3.txt').read_text(encoding='utf8').splitlines():
    if not line.strip():
        continue
    if line.startswith('# '):
        para(line[2:], 'Title')
    elif line == '!AUTHORS':
        for b in orig[1:4]:
            append_block(b)
    elif line.startswith('### '):
        para(line[4:], 'Heading 2')
    elif line.startswith('## '):
        para(line[3:], 'Heading 1')
    elif line.startswith('!EQ'):
        equations(int(line[3:]))
    elif line.startswith('!FIG'):
        token, cap = line.split('|', 1)
        n = int(token[4:])
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.keep_with_next = True
        p.paragraph_format.space_after = Pt(3)
        p.add_run().add_picture(str(F / f'figure{n}.png'), width=Inches(6.45))
        caption(f'Figure {n}. ' + cap)
    elif line.startswith('!TABLE'):
        tables(line[6:])
    elif line == '!BIBLIOGRAPHY':
        for b in bib:
            append_block(b)
    elif line == '!NEWREFERENCES':
        for ref in refs:
            p = doc.add_paragraph()
            add_plain(p, ref, True)
            p.paragraph_format.space_after = Pt(6)
            p.paragraph_format.line_spacing = 1.0
    else:
        para(line)

assert sorted(used) == list(range(30)), sorted(used)
doc.save(DST)

# ---- verification ----------------------------------------------------------
with ZipFile(SRC_DOC) as za, ZipFile(DST) as zb:
    xa = etree.fromstring(za.read('word/document.xml'))
    xb = etree.fromstring(zb.read('word/document.xml'))
    ia = xa.xpath('//w:instrText/text()', namespaces=NS)
    ib = xb.xpath('//w:instrText/text()', namespaces=NS)
    assert Counter(ia) == Counter(ib), 'Zotero instruction content changed'
    custom_same = za.read('docProps/custom.xml') == zb.read('docProps/custom.xml')
    assert custom_same
    ta = xa.xpath('//w:fldChar/@w:fldCharType', namespaces=NS)
    tb = xb.xpath('//w:fldChar/@w:fldCharType', namespaces=NS)
    assert Counter(ta) == Counter(tb)
    sa, sb = xa.find('w:body', NS), xb.find('w:body', NS)
    btxt = [''.join(e.xpath('.//w:t/text()', namespaces=NS)) for e in list(sa)[60:89]]
    alltxt = [''.join(e.xpath('.//w:t/text()', namespaces=NS)) for e in sb]
    assert all(t in alltxt for t in btxt)
    report = dict(
        output=str(DST), source_sha256=source_hash,
        source_unchanged=hashlib.sha256(SRC_DOC.read_bytes()).hexdigest() == source_hash,
        zotero_instruction_nodes=len(ia),
        zotero_instruction_content_identical=True,
        original_citation_fields_preserved=30,
        zotero_custom_properties_byte_identical=custom_same,
        bibliography_cached_text_preserved=True,
        native_math_objects=len(xb.findall('.//' + qn('m:oMath'))),
        embedded_images=len(xb.findall('.//' + qn('w:drawing'))),
        additional_references=len(refs),
        word_count=len(' '.join(xb.xpath('//w:t/text()', namespaces=NS)).split()))
    (W / 'document_verification_v3.json').write_text(json.dumps(report, indent=2))
    print(json.dumps(report, indent=2))
